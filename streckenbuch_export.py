#!/usr/bin/env python3
"""Betriebsstellen und Rufnummern aus einem DB-Streckenbuch exportieren.

Der Parser ist absichtlich heuristisch: PDF-Text besitzt nicht immer dieselbe
Reihenfolge wie die sichtbar gesetzten Tabellen. Jede Zeile behält deshalb ihre
Fundstelle und eine Konfidenzbewertung für die fachliche Nachkontrolle.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence


PROGRAM_NAME = "Streckenbuch Export"
PROGRAM_VERSION = "0.3.0"


PRIVATE_USE_RE = re.compile(r"[\ue000-\uf8ff]")
SPACE_RE = re.compile(r"\s+")
ROUTE_RE = re.compile(
    r"^(?P<number>\d{1,4}(?:[.,]\d+)?(?:\s+[A-Z])?)\s+"
    r"(?P<title>(?!Regeln\b|Ril\b|Hinweise\b).{4,})$",
    re.IGNORECASE,
)
ROUTE_ONLY_RE = re.compile(r"^\d{1,4}(?:[.,]\d+)?(?:\s+[A-Z])?$")
STATION_RE = re.compile(
    r"^(?P<type>Bf|Bft|Bfu|Bbf|Abzw|Awanst|Anst|Hp|Hst|Üst|Bk|Bftst)\s+"
    r"(?P<name>.+?)$",
    re.IGNORECASE,
)
RULE_RE = re.compile(r"^(Ril\s+\d{3}(?:\.\d{4})?.*)$", re.IGNORECASE)
SHORT_LABELED_RE = re.compile(r"\bKurzwahl\s*[:=]?\s*(\d{4})\b", re.IGNORECASE)
LONG_NUMBER_RE = re.compile(r"(?<!\d)(7(?:[\s\u00a0]*\d){7})(?!\d)")
NINE_DIGIT_RE = re.compile(r"(?<!\d)(9\d{8})(?!\d)")
TABLE_PAIR_RE = re.compile(
    r"(?<!\d)(\d{4})\s*[,;]?\s+(7(?:[\s\u00a0]*\d){7})(?!\d)"
)
TCOM_RE = re.compile(
    r"\bT-?Com\s*:\s*(\+?\d[\d\s/()\-]{5,}\d)", re.IGNORECASE
)
ROLE_RE = re.compile(
    r"\b(Fdl(?:-Helfer)?|Ww|BÜ|Disponent|Anlagendisponent|BöZM|Zm|özF)\b",
    re.IGNORECASE,
)
SENTENCE_WORDS_RE = re.compile(
    r"\b(ist|sind|wird|werden|muss|müssen|darf|dürfen|siehe|zwischen|zugelassen|umfasst|befindet)\b",
    re.IGNORECASE,
)
TECHNICAL_STATION_FALSE_POSITIVE_RE = re.compile(
    r"\b(Gleise?|BÜ|[AEZ]sig|Ls|Weiche|km/h)\b|‰", re.IGNORECASE
)


@dataclass(slots=True)
class ParseContext:
    route_number: str = ""
    route_name: str = ""
    station_type: str = ""
    station_name: str = ""
    rule: str = ""
    in_station_section: bool = False


@dataclass(slots=True)
class ContactRecord:
    route_number: str = ""
    route_name: str = ""
    station_type: str = ""
    station_name: str = ""
    role: str = ""
    position: str = ""
    short_number: str = ""
    long_number: str = ""
    gsm_r_number: str = ""
    tcom_number: str = ""
    functional_number: str = ""
    number_type: str = ""
    rule: str = ""
    pdf_page: int = 0
    document_page: int | None = None
    source_kind: str = ""
    confidence: float = 0.0
    source_text: str = ""


@dataclass(slots=True, frozen=True)
class StationRecord:
    route_number: str
    route_name: str
    station_type: str
    station_name: str
    pdf_page: int
    document_page: int | None


@dataclass(slots=True)
class PageParseResult:
    context: ParseContext
    records: list[ContactRecord] = field(default_factory=list)
    stations: list[StationRecord] = field(default_factory=list)


CONTACT_COLUMNS = [
    ("route_number", "Strecke"),
    ("route_name", "Streckenbezeichnung"),
    ("station_type", "Betriebsstellenart"),
    ("station_name", "Betriebsstelle"),
    ("role", "Rolle"),
    ("position", "Stelle/Funktion"),
    ("short_number", "Kurzwahl"),
    ("long_number", "Langwahl"),
    ("gsm_r_number", "GSM-R"),
    ("tcom_number", "T-Com"),
    ("functional_number", "Funktionale Rufnummer"),
    ("number_type", "Rufnummerntyp"),
    ("rule", "Ril/Regel"),
    ("pdf_page", "PDF-Seite"),
    ("document_page", "Dokumentseite"),
    ("source_kind", "Fundstellenart"),
    ("confidence", "Konfidenz"),
    ("source_text", "Originaltext"),
]


def normalize_line(line: str) -> str:
    """Private Symbolfonts entfernen und Leerraum vereinheitlichen."""
    line = unicodedata.normalize("NFKC", line)
    line = PRIVATE_USE_RE.sub(" ", line)
    line = line.replace("–", "-").replace("—", "-").replace("−", "-")
    return SPACE_RE.sub(" ", line).strip()


def digits(value: str) -> str:
    return re.sub(r"\D", "", value)


def printed_page(lines: Sequence[str]) -> int | None:
    for line in reversed(lines[-8:]):
        if re.fullmatch(r"\d{1,4}", line):
            return int(line)
    return None


def route_from_lines(lines: Sequence[str]) -> tuple[str, str] | None:
    """Die Streckenüberschrift im oberen Seitenbereich erkennen.

    Auf Folgeseiten steht oben nur die Streckennummer. Ohne die Prüfung gegen
    diese Kopfzeile könnten Tabellenzeilen wie ``831 bis Gleisabschluss ...``
    fälschlich als neue Strecke interpretiert werden.
    """
    if not lines:
        return None

    first_match = ROUTE_RE.match(lines[0])
    candidates: list[re.Match[str]] = [first_match] if first_match else []

    for header_index, header_number in enumerate(lines[:4]):
        if not ROUTE_ONLY_RE.fullmatch(header_number):
            continue
        normalized_header = header_number.replace(",", ".")
        for line in lines[header_index + 1 : header_index + 5]:
            match = ROUTE_RE.match(line)
            if match and match.group("number").replace(",", ".") == normalized_header:
                candidates.append(match)

    for match in candidates:
        title = match.group("title").strip(" -*↕↑↓→←")
        if "Streckenverzeichnis" in title or title.isdigit():
            continue
        return match.group("number").replace(",", "."), title
    return None


def station_from_line(line: str) -> tuple[str, str] | None:
    candidate = line.lstrip("↓↑→←↕ ")
    match = STATION_RE.match(candidate)
    if not match or len(candidate) > 180:
        return None
    raw_name = LONG_NUMBER_RE.split(match.group("name"), maxsplit=1)[0]
    raw_name = re.sub(r"\s*[☎☏]+\s*$", "", raw_name).strip(" ,;:-*")
    if (
        not raw_name
        or SENTENCE_WORDS_RE.search(raw_name)
        or TECHNICAL_STATION_FALSE_POSITIVE_RE.search(raw_name)
    ):
        return None

    # Zusatzangaben beschreiben oft Fernsteuerung, Zuständigkeit oder nur den
    # folgenden Unterabschnitt. Für die Suchspalte ist der stabile Name der
    # Betriebsstelle hilfreicher; der vollständige Text bleibt in Originaltext.
    name = raw_name.split(",", 1)[0]
    name = re.sub(
        r"\s+\((?:özF|ferngesteuert|ferngestellt|Fdl)\b.*\)\s*$",
        "",
        name,
        flags=re.IGNORECASE,
    )
    if LONG_NUMBER_RE.search(line):
        name = re.split(r"\s+(?:Fdl|Ww)\b", name, maxsplit=1, flags=re.IGNORECASE)[0]
    name = name.strip(" ,;:-*")
    return (match.group("type"), name) if name else None


def extract_position(line: str, number_spans: Sequence[tuple[int, int]]) -> str:
    if not number_spans:
        return ""
    first_number = min(start for start, _ in number_spans)
    prefix = line[:first_number]
    prefix = re.sub(r"\b(?:Kurz|Lang)wahl\s*[:=]?\s*$", "", prefix, flags=re.I)
    prefix = re.sub(r"\b(?:GSM-R|T-?Com|Rufnummer)\s*[:=]?\s*$", "", prefix, flags=re.I)
    prefix = prefix.strip(" ,;:-")
    if prefix.lower() in {
        "stelle",
        "kurzwahl",
        "langwahl",
        "erreichbarkeit",
        "rufnummer",
        "rufnummern",
    }:
        return ""
    return prefix[-140:]


def detect_role(position: str) -> str:
    match = ROLE_RE.search(position)
    return match.group(1) if match else ""


def classify_types(record: ContactRecord) -> str:
    types: list[str] = []
    if record.short_number:
        types.append("Kurzwahl")
    if record.long_number:
        types.append("Langwahl")
    if record.gsm_r_number:
        types.append("GSM-R")
    if record.tcom_number:
        types.append("T-Com")
    if record.functional_number:
        types.append("funktional")
    return ", ".join(types)


def make_record(
    context: ParseContext,
    *,
    pdf_page: int,
    document_page: int | None,
    line: str,
    position: str,
    short: str = "",
    long: str = "",
    gsm_r: str = "",
    tcom: str = "",
    functional: str = "",
    source_kind: str,
    confidence: float,
) -> ContactRecord:
    if functional == "999010004" and not position:
        position = "Zugvorbereitungsmeldung"
    record = ContactRecord(
        route_number=context.route_number,
        route_name=context.route_name,
        station_type=context.station_type,
        station_name=context.station_name,
        role=detect_role(position),
        position=position,
        short_number=short,
        long_number=long,
        gsm_r_number=gsm_r,
        tcom_number=tcom,
        functional_number=functional,
        rule=context.rule,
        pdf_page=pdf_page,
        document_page=document_page,
        source_kind=source_kind,
        confidence=confidence,
        source_text=line,
    )
    record.number_type = classify_types(record)
    return record


def records_from_line(
    line: str,
    context: ParseContext,
    *,
    pdf_page: int,
    document_page: int | None,
) -> list[ContactRecord]:
    """Rufnummern einer Textzeile erkennen und soweit möglich zusammenführen."""
    short_matches = list(SHORT_LABELED_RE.finditer(line))
    long_matches = list(LONG_NUMBER_RE.finditer(line))
    nine_matches = list(NINE_DIGIT_RE.finditer(line))
    tcom_matches = list(TCOM_RE.finditer(line))

    table_pairs = list(TABLE_PAIR_RE.finditer(line))
    known_shorts = {m.group(1) for m in short_matches}
    for match in table_pairs:
        if match.group(1) not in known_shorts:
            # Ein minimales Match-ähnliches Objekt wird nicht benötigt; die
            # Werte und Spans werden separat weitergereicht.
            known_shorts.add(match.group(1))

    shorts = [m.group(1) for m in short_matches]
    for match in table_pairs:
        if match.group(1) not in shorts:
            shorts.append(match.group(1))
    longs = [digits(m.group(1)) for m in long_matches]

    spans = [m.span() for m in short_matches + long_matches + nine_matches + tcom_matches]
    spans.extend(m.span() for m in table_pairs)
    position = extract_position(line, spans)
    records: list[ContactRecord] = []

    if longs:
        for index, long_number in enumerate(longs):
            short_number = ""
            if len(shorts) == 1:
                short_number = shorts[0]
            elif index < len(shorts):
                short_number = shorts[index]
            labeled = bool(short_matches) or "Langwahl" in line
            confidence = 0.96 if labeled and short_number else 0.88 if labeled else 0.72
            records.append(
                make_record(
                    context,
                    pdf_page=pdf_page,
                    document_page=document_page,
                    line=line,
                    position=position,
                    short=short_number,
                    long=long_number,
                    source_kind="Kurzwahl/Langwahl" if short_number else "Langwahl/Fundstelle",
                    confidence=confidence,
                )
            )
    elif shorts:
        for short_number in shorts:
            records.append(
                make_record(
                    context,
                    pdf_page=pdf_page,
                    document_page=document_page,
                    line=line,
                    position=position,
                    short=short_number,
                    source_kind="Kurzwahl/Fundstelle",
                    confidence=0.75,
                )
            )

    nine_numbers = [m.group(1) for m in nine_matches]
    tcom_numbers = [SPACE_RE.sub(" ", m.group(1)).strip() for m in tcom_matches]

    # Eine einzelne Kontaktzeile kann Kurz-/Langwahl, GSM-R und T-Com gemeinsam
    # enthalten. In diesem Fall bleiben alle Werte in einer Excel-Zeile.
    if len(records) == 1:
        primary = records[0]
        for number in nine_numbers:
            if number == "999010004" or "Zugvorbereitung" in line:
                primary.functional_number = number
            elif "GSM-R" in line:
                primary.gsm_r_number = number
            else:
                primary.functional_number = number
        if tcom_numbers:
            primary.tcom_number = tcom_numbers[0]
        primary.number_type = classify_types(primary)
    else:
        for number in nine_numbers:
            is_functional = number == "999010004" or "Zugvorbereitung" in line
            records.append(
                make_record(
                    context,
                    pdf_page=pdf_page,
                    document_page=document_page,
                    line=line,
                    position=position,
                    functional=number if is_functional else "",
                    gsm_r="" if is_functional else number,
                    source_kind="Funktionale Rufnummer" if is_functional else "GSM-R/Fundstelle",
                    confidence=0.98 if is_functional else 0.9,
                )
            )
        for number in tcom_numbers:
            records.append(
                make_record(
                    context,
                    pdf_page=pdf_page,
                    document_page=document_page,
                    line=line,
                    position=position,
                    tcom=number,
                    source_kind="T-Com/Fundstelle",
                    confidence=0.9,
                )
            )

    return records


def parse_page(
    text: str,
    *,
    pdf_page: int,
    previous_context: ParseContext | None = None,
) -> PageParseResult:
    lines = [normalize_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    context = replace(previous_context) if previous_context else ParseContext()
    page_number = printed_page(lines)

    route = route_from_lines(lines)
    if route and route != (context.route_number, context.route_name):
        context = ParseContext(route_number=route[0], route_name=route[1])

    records: list[ContactRecord] = []
    stations: list[StationRecord] = []

    for line in lines:
        if "2. Regeln für Betriebsstellen" in line:
            context.in_station_section = True
            continue
        if "1. Regeln für die Strecke" in line:
            context.in_station_section = False
            context.station_type = ""
            context.station_name = ""
            context.rule = ""
            continue

        rule_match = RULE_RE.match(line)
        if rule_match:
            context.rule = rule_match.group(1)

        if context.in_station_section:
            station = station_from_line(line)
            if station:
                context.station_type, context.station_name = station
                stations.append(
                    StationRecord(
                        route_number=context.route_number,
                        route_name=context.route_name,
                        station_type=context.station_type,
                        station_name=context.station_name,
                        pdf_page=pdf_page,
                        document_page=page_number,
                    )
                )

        # Rufnummern vor Abschnitt 2 sind selten, aber funktionale Nummern und
        # klar beschriftete Kontaktzeilen sollen trotzdem auffindbar bleiben.
        if (
            context.in_station_section
            or "Kurzwahl" in line
            or "Langwahl" in line
            or "GSM-R" in line
            or "T-Com" in line
            or "Rufnummer" in line
        ):
            records.extend(
                records_from_line(
                    line,
                    context,
                    pdf_page=pdf_page,
                    document_page=page_number,
                )
            )

    return PageParseResult(context=context, records=deduplicate(records), stations=stations)


def deduplicate(records: Iterable[ContactRecord]) -> list[ContactRecord]:
    """Identische Parser-Treffer innerhalb derselben PDF-Seite entfernen."""
    seen: set[tuple[object, ...]] = set()
    result: list[ContactRecord] = []
    for record in records:
        key = contact_duplicate_key(record) + (record.pdf_page,)
        if key in seen:
            continue
        seen.add(key)
        result.append(record)
    return result


def normalized_key_part(value: str) -> str:
    return normalize_line(value).strip(" ,;:-").casefold()


def contact_duplicate_key(record: ContactRecord) -> tuple[str, ...]:
    """Fachlichen Kontaktschlüssel ohne Fundseite und Quelltext bilden."""
    return (
        normalized_key_part(record.route_number),
        normalized_key_part(record.station_type),
        normalized_key_part(record.station_name),
        normalized_key_part(record.position),
        record.short_number,
        record.long_number,
        record.gsm_r_number,
        normalized_key_part(record.tcom_number),
        record.functional_number,
    )


def deduplicate_for_export(records: Iterable[ContactRecord]) -> list[ContactRecord]:
    """Wiederholte Kontakte seitenübergreifend vor dem Export entfernen.

    Strecke, Betriebsstelle, Funktion und sämtliche Rufnummern gehören zur
    Identität. Ril, Seite, Konfidenz und Originaltext sind Provenienz des ersten
    Treffers und beeinflussen die Dublettenerkennung nicht.
    """
    seen: set[tuple[str, ...]] = set()
    result: list[ContactRecord] = []
    for record in records:
        key = contact_duplicate_key(record)
        if key in seen:
            continue
        seen.add(key)
        result.append(record)
    return result


def extract_pdf(
    pdf_path: Path, selected_pages: Sequence[int] | None = None
) -> tuple[list[ContactRecord], list[StationRecord], int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - CLI-Hinweis
        raise SystemExit("pypdf fehlt. Installation: uv sync") from exc

    reader = PdfReader(str(pdf_path), strict=False)
    page_indexes = list(selected_pages) if selected_pages is not None else list(range(len(reader.pages)))
    context = ParseContext()
    records: list[ContactRecord] = []
    stations: list[StationRecord] = []

    for count, page_index in enumerate(page_indexes, start=1):
        page = reader.pages[page_index]
        try:
            text = page.extract_text(extraction_mode="layout") or ""
        except (TypeError, ValueError):
            text = page.extract_text() or ""
        result = parse_page(text, pdf_page=page_index + 1, previous_context=context)
        context = result.context
        records.extend(result.records)
        stations.extend(result.stations)
        if count % 50 == 0 or count == len(page_indexes):
            print(
                f"Verarbeitet: {count}/{len(page_indexes)} PDF-Seiten, "
                f"{len(records)} Rufnummern-Fundstellen",
                file=sys.stderr,
            )

    return deduplicate(records), deduplicate_stations(stations), len(reader.pages)


def deduplicate_stations(stations: Iterable[StationRecord]) -> list[StationRecord]:
    seen: set[tuple[str, str, str]] = set()
    result: list[StationRecord] = []
    for station in stations:
        key = (station.route_number, station.station_type, station.station_name)
        if key in seen:
            continue
        seen.add(key)
        result.append(station)
    return result


def parse_page_selection(spec: str, page_count: int) -> list[int]:
    selected: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start_text, end_text = part.split("-", 1)
            start, end = int(start_text), int(end_text)
            if end < start:
                raise ValueError(f"Ungültiger Seitenbereich: {part}")
            selected.update(range(start - 1, end))
        else:
            selected.add(int(part) - 1)
    if not selected or min(selected) < 0 or max(selected) >= page_count:
        raise ValueError(f"Seiten müssen zwischen 1 und {page_count} liegen")
    return sorted(selected)


def write_csv(records: Sequence[ContactRecord], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([label for _, label in CONTACT_COLUMNS])
        for record in records:
            values = asdict(record)
            writer.writerow([values[field_name] for field_name, _ in CONTACT_COLUMNS])


def write_xlsx(
    records: Sequence[ContactRecord],
    stations: Sequence[StationRecord],
    path: Path,
    *,
    source_pdf: Path,
    total_pages: int,
    duplicates_removed: int = 0,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:  # pragma: no cover - CLI-Hinweis
        raise SystemExit("openpyxl fehlt. Installation: uv sync") from exc

    workbook = Workbook()
    contacts_sheet = workbook.active
    contacts_sheet.title = "Kontakte"
    headers = [label for _, label in CONTACT_COLUMNS]
    contacts_sheet.append(headers)

    for record in records:
        values = asdict(record)
        contacts_sheet.append(
            [values[field_name] for field_name, _ in CONTACT_COLUMNS]
        )

    station_sheet = workbook.create_sheet("Betriebsstellen")
    station_headers = [
        "Strecke",
        "Streckenbezeichnung",
        "Betriebsstellenart",
        "Betriebsstelle",
        "Erste PDF-Seite",
        "Dokumentseite",
    ]
    station_sheet.append(station_headers)
    for station in stations:
        station_sheet.append(
            [
                station.route_number,
                station.route_name,
                station.station_type,
                station.station_name,
                station.pdf_page,
                station.document_page,
            ]
        )

    info_sheet = workbook.create_sheet("Info")
    info_rows = [
        ("Quelldatei", str(source_pdf)),
        ("PDF-Seiten gesamt", total_pages),
        ("Rufnummern-Fundstellen nach Dublettenbereinigung", len(records)),
        ("Entfernte Dubletten", duplicates_removed),
        ("Eindeutige Betriebsstellen je Strecke", len(stations)),
        ("Erstellt am", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Parser-Version", PROGRAM_VERSION),
        (
            "Hinweis",
            "Heuristischer Prototyp; Fundstellen vor betrieblicher Nutzung im gültigen Quelldokument prüfen.",
        ),
    ]
    for row in info_rows:
        info_sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet, widths in (
        (
            contacts_sheet,
            [12, 38, 19, 30, 16, 38, 12, 14, 14, 20, 21, 24, 24, 11, 15, 22, 12, 70],
        ),
        (station_sheet, [12, 42, 19, 38, 16, 16]),
    ):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        if sheet.max_row >= 2:
            table_name = "KontakteTabelle" if sheet is contacts_sheet else "BetriebsstellenTabelle"
            table = Table(displayName=table_name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

    # Rufnummernspalten als Text formatieren.
    for column in range(7, 12):
        for row in range(2, contacts_sheet.max_row + 1):
            contacts_sheet.cell(row=row, column=column).number_format = "@"

    info_sheet.column_dimensions["A"].width = 48
    info_sheet.column_dimensions["B"].width = 100
    info_sheet["A1"].font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def spreadsheet_column_name(index: int) -> str:
    """1-basierten Spaltenindex in Calc-/Excel-Spaltennamen umwandeln."""
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def write_ods(
    records: Sequence[ContactRecord],
    stations: Sequence[StationRecord],
    path: Path,
    *,
    source_pdf: Path,
    total_pages: int,
    duplicates_removed: int = 0,
) -> None:
    """Native LibreOffice-Calc-Datei mit Filterfeldern erzeugen."""
    try:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.style import (
            Style,
            TableCellProperties,
            TableColumnProperties,
            TextProperties,
        )
        from odf.table import (
            DatabaseRange,
            DatabaseRanges,
            Table,
            TableCell,
            TableColumn,
            TableHeaderRows,
            TableRow,
        )
        from odf.text import P
    except ImportError as exc:  # pragma: no cover - CLI-Hinweis
        raise SystemExit("odfpy fehlt. Installation: uv sync") from exc

    document = OpenDocumentSpreadsheet()

    header_style = Style(name="HeaderCell", family="table-cell")
    header_style.addElement(
        TableCellProperties(backgroundcolor="#1F4E78", padding="0.08in")
    )
    header_style.addElement(TextProperties(color="#FFFFFF", fontweight="bold"))
    document.styles.addElement(header_style)

    body_style = Style(name="BodyCell", family="table-cell")
    body_style.addElement(
        TableCellProperties(
            border="0.01pt solid #D9E2F3",
            padding="0.04in",
            verticalalign="top",
            wrapoption="wrap",
        )
    )
    document.styles.addElement(body_style)

    def add_columns(table: Table, widths: Sequence[int], prefix: str) -> None:
        for index, width in enumerate(widths, start=1):
            column_style = Style(name=f"{prefix}Column{index}", family="table-column")
            column_style.addElement(
                TableColumnProperties(columnwidth=f"{min(max(width * 0.18, 1.8), 13):.2f}cm")
            )
            document.automaticstyles.addElement(column_style)
            table.addElement(TableColumn(stylename=column_style))

    def make_cell(value: object, *, header: bool = False, numeric: bool = False) -> TableCell:
        style = header_style if header else body_style
        if numeric and value not in (None, ""):
            cell = TableCell(valuetype="float", value=float(value), stylename=style)
        else:
            cell = TableCell(valuetype="string", stylename=style)
        cell.addElement(P(text="" if value is None else str(value)))
        return cell

    def add_header(table: Table, labels: Sequence[str]) -> None:
        header_rows = TableHeaderRows()
        row = TableRow()
        for label in labels:
            row.addElement(make_cell(label, header=True))
        header_rows.addElement(row)
        table.addElement(header_rows)

    contacts_table = Table(name="Kontakte")
    contact_widths = [12, 38, 19, 30, 16, 38, 12, 14, 14, 20, 21, 24, 24, 11, 15, 22, 12, 70]
    add_columns(contacts_table, contact_widths, "Contacts")
    add_header(contacts_table, [label for _, label in CONTACT_COLUMNS])
    numeric_contact_fields = {"pdf_page", "document_page", "confidence"}
    for record in records:
        values = asdict(record)
        row = TableRow()
        for field_name, _ in CONTACT_COLUMNS:
            row.addElement(
                make_cell(
                    values[field_name],
                    numeric=field_name in numeric_contact_fields,
                )
            )
        contacts_table.addElement(row)
    document.spreadsheet.addElement(contacts_table)

    station_headers = [
        "Strecke",
        "Streckenbezeichnung",
        "Betriebsstellenart",
        "Betriebsstelle",
        "Erste PDF-Seite",
        "Dokumentseite",
    ]
    stations_table = Table(name="Betriebsstellen")
    add_columns(stations_table, [12, 42, 19, 38, 16, 16], "Stations")
    add_header(stations_table, station_headers)
    for station in stations:
        row = TableRow()
        values = [
            station.route_number,
            station.route_name,
            station.station_type,
            station.station_name,
            station.pdf_page,
            station.document_page,
        ]
        for index, value in enumerate(values):
            row.addElement(make_cell(value, numeric=index >= 4))
        stations_table.addElement(row)
    document.spreadsheet.addElement(stations_table)

    info_table = Table(name="Info")
    add_columns(info_table, [48, 100], "Info")
    info_rows = [
        ("Quelldatei", str(source_pdf)),
        ("PDF-Seiten gesamt", total_pages),
        ("Rufnummern-Fundstellen nach Dublettenbereinigung", len(records)),
        ("Entfernte Dubletten", duplicates_removed),
        ("Eindeutige Betriebsstellen je Strecke", len(stations)),
        ("Erstellt am", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Parser-Version", PROGRAM_VERSION),
        (
            "Hinweis",
            "Heuristischer Prototyp; Fundstellen vor betrieblicher Nutzung im gültigen Quelldokument prüfen.",
        ),
    ]
    for label, value in info_rows:
        row = TableRow()
        row.addElement(make_cell(label, header=True))
        row.addElement(make_cell(value, numeric=isinstance(value, int)))
        info_table.addElement(row)
    document.spreadsheet.addElement(info_table)

    # Calc zeigt für native ODS-Dateien Filter-Dropdowns nur dann automatisch
    # an, wenn ein table:database-range mit display-filter-buttons gesetzt ist.
    database_ranges = DatabaseRanges()
    contact_last_column = spreadsheet_column_name(len(CONTACT_COLUMNS))
    database_ranges.addElement(
        DatabaseRange(
            name="KontakteFilter",
            targetrangeaddress=(
                f"Kontakte.A1:Kontakte.{contact_last_column}{max(1, len(records) + 1)}"
            ),
            containsheader=True,
            displayfilterbuttons=True,
        )
    )
    database_ranges.addElement(
        DatabaseRange(
            name="BetriebsstellenFilter",
            targetrangeaddress=(
                f"Betriebsstellen.A1:Betriebsstellen.F{max(1, len(stations) + 1)}"
            ),
            containsheader=True,
            displayfilterbuttons=True,
        )
    )
    document.spreadsheet.addElement(database_ranges)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path), addsuffix=False)


def output_paths(
    input_pdf: Path,
    output: Path | None,
    output_format: str,
) -> tuple[Path | None, Path | None, Path | None]:
    base = output if output else input_pdf.with_name(f"{input_pdf.stem}_rufnummern.xlsx")
    if base.suffix.lower() in {".xlsx", ".ods", ".csv"}:
        stem = base.with_suffix("")
    else:
        stem = base
    xlsx_path = stem.with_suffix(".xlsx") if output_format in {"xlsx", "both", "all"} else None
    ods_path = stem.with_suffix(".ods") if output_format in {"ods", "all"} else None
    csv_path = stem.with_suffix(".csv") if output_format in {"csv", "both", "all"} else None
    return xlsx_path, ods_path, csv_path


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Betriebsstellen und Rufnummern aus einem Streckenbuch nach "
            "Excel, LibreOffice Calc oder CSV exportieren."
        )
    )
    parser.add_argument("pdf", type=Path, nargs="?", help="Eingabe-PDF")
    parser.add_argument(
        "--version",
        action="version",
        version=f"{PROGRAM_NAME} {PROGRAM_VERSION}",
    )
    parser.add_argument("-o", "--output", type=Path, help="Ausgabedatei bzw. Ausgabestamm")
    parser.add_argument(
        "--format",
        choices=("xlsx", "ods", "csv", "both", "all"),
        default="xlsx",
        help=(
            "Ausgabeformat: ods ist nativ für LibreOffice, both erzeugt "
            "XLSX+CSV, all erzeugt XLSX+ODS+CSV (Standard: xlsx)"
        ),
    )
    parser.add_argument(
        "--pages",
        help="Optionale PDF-Seiten, z. B. 15-30,42 (1-basiert)",
    )
    parser.add_argument(
        "--min-confidence",
        type=float,
        default=0.0,
        help="Nur Fundstellen ab dieser Konfidenz exportieren (0.0-1.0)",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_argument_parser()
    args = parser.parse_args(argv)
    if args.pdf is None:
        parser.error("Eingabe-PDF fehlt")
    pdf_path: Path = args.pdf.expanduser().resolve()
    if not pdf_path.is_file():
        print(f"PDF nicht gefunden: {pdf_path}", file=sys.stderr)
        return 2
    if not 0 <= args.min_confidence <= 1:
        print("--min-confidence muss zwischen 0 und 1 liegen", file=sys.stderr)
        return 2

    selected_pages: list[int] | None = None
    if args.pages:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise SystemExit("pypdf fehlt. Installation: uv sync") from exc
        page_count = len(PdfReader(str(pdf_path), strict=False).pages)
        try:
            selected_pages = parse_page_selection(args.pages, page_count)
        except ValueError as exc:
            print(str(exc), file=sys.stderr)
            return 2

    records, stations, total_pages = extract_pdf(pdf_path, selected_pages)
    records = [record for record in records if record.confidence >= args.min_confidence]
    records_before_deduplication = len(records)
    records = deduplicate_for_export(records)
    duplicates_removed = records_before_deduplication - len(records)
    xlsx_path, ods_path, csv_path = output_paths(
        pdf_path,
        args.output.expanduser().resolve() if args.output else None,
        args.format,
    )
    if xlsx_path:
        write_xlsx(
            records,
            stations,
            xlsx_path,
            source_pdf=pdf_path,
            total_pages=total_pages,
            duplicates_removed=duplicates_removed,
        )
        print(f"Excel:       {xlsx_path}")
    if ods_path:
        write_ods(
            records,
            stations,
            ods_path,
            source_pdf=pdf_path,
            total_pages=total_pages,
            duplicates_removed=duplicates_removed,
        )
        print(f"LibreOffice: {ods_path}")
    if csv_path:
        write_csv(records, csv_path)
        print(f"CSV:         {csv_path}")
    print(
        f"Exportiert: {len(records)} eindeutige Fundstellen, "
        f"{duplicates_removed} Dubletten entfernt, "
        f"{len(stations)} Betriebsstellen"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
